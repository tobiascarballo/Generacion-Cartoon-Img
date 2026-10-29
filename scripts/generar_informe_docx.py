#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
============================================================================
Trabajo Practico Final Integrador - Computacion Avanzada (UADER FCyT)
Alumnos: Carballo Tobias, Dacuy Santino

Script: generar_informe_docx.py
Descripcion:
  Genera el informe oficial en formato Microsoft Word (.docx) respetando el
  membrete institucional de UADER FCyT (extraido de TPs/Ejercicio7.docx),
  incluyendo todo el marco teorico, explicacion algoritmica, tablas de
  Speedup y Eficiencia, y los 6 graficos/paneles visuales embebidos.
============================================================================
"""

import os
import csv
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def extraer_logo_uader(base_dir):
    docx_ref = os.path.join(os.path.dirname(base_dir), "Ejercicio7.docx")
    logo_path = os.path.join(base_dir, "docs", "uader_fcyt_logo.jpg")
    os.makedirs(os.path.dirname(logo_path), exist_ok=True)
    if os.path.exists(docx_ref):
        with zipfile.ZipFile(docx_ref, "r") as z:
            if "word/media/image1.jpg" in z.namelist():
                with open(logo_path, "wb") as f:
                    f.write(z.read("word/media/image1.jpg"))
                return logo_path
    return None


def agregar_titulo_seccion(doc, texto, nivel=1):
    h = doc.add_heading(texto, level=nivel)
    for run in h.runs:
        run.font.color.rgb = RGBColor(15, 55, 110)
    return h


def cargar_benchmarks(csv_path):
    filas = []
    with open(csv_path, mode="r", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            r["procesos"] = int(r["procesos"])
            r["threads"] = int(r["threads"])
            r["workers"] = int(r["workers"])
            r["rangos"] = int(r["rangos"])
            r["tiempo_total"] = float(r["tiempo_total"])
            r["tiempo_computo"] = float(r["tiempo_computo"])
            r["speedup"] = float(r["speedup"])
            r["eficiencia"] = float(r["eficiencia"])
            filas.append(r)
    return filas


def agregar_tabla_resolucion(doc, filas, res_str):
    configs = [
        ("SEQ", 1, 1, "Secuencial (1x1)"),
        ("OMP", 1, 2, "OpenMP (2 threads)"),
        ("OMP", 1, 4, "OpenMP (4 threads)"),
        ("OMP", 1, 8, "OpenMP (8 threads)"),
        ("OMP", 1, 16, "OpenMP (16 threads)"),
        ("MPI", 2, 1, "MPI (2 procesos)"),
        ("MPI", 4, 1, "MPI (4 procesos)"),
        ("MPI", 8, 1, "MPI (8 procesos)"),
        ("HYBRID", 2, 2, "Híbrido (2P x 2T)"),
        ("HYBRID", 2, 4, "Híbrido (2P x 4T)"),
        ("HYBRID", 4, 2, "Híbrido (4P x 2T)"),
        ("HYBRID", 4, 4, "Híbrido (4P x 4T)"),
    ]

    tabla = doc.add_table(rows=1, cols=6)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    encabezados = [
        "Configuración",
        "Workers (p)",
        "Filtro 3x3 | Post 3\n(T / S / E)",
        "Filtro 3x3 | Post 9\n(T / S / E)",
        "Filtro 5x5 | Post 3\n(T / S / E)",
        "Filtro 5x5 | Post 9\n(T / S / E)",
    ]
    hdr_cells = tabla.rows[0].cells
    for i, txt in enumerate(encabezados):
        hdr_cells[i].text = txt
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(8.5)

    combos = [("3x3", 3), ("3x3", 9), ("5x5", 3), ("5x5", 9)]

    for par, pr, th, nombre_cfg in configs:
        row_cells = tabla.add_row().cells
        row_cells[0].text = nombre_cfg
        row_cells[1].text = str(pr * th)
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for idx_c, (flt, rng) in enumerate(combos):
            match = [
                r for r in filas
                if r["resolucion"] == res_str
                and r["paradigma"] == par
                and r["procesos"] == pr
                and r["threads"] == th
                and r["filtro"] == flt
                and r["rangos"] == rng
            ]
            if match:
                m = match[0]
                celda_txt = f"{m['tiempo_total']:.4f}s / {m['speedup']:.2f}x / {m['eficiencia']*100:.0f}%"
            else:
                celda_txt = "-"
            row_cells[2 + idx_c].text = celda_txt
            row_cells[2 + idx_c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for c in row_cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)

    doc.add_paragraph()


def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    csv_path = os.path.join(base_dir, "results", "benchmarks.csv")
    plots_dir = os.path.join(base_dir, "results", "plots")
    out_docx = os.path.join(base_dir, "docs", "Informe_TP_Final_Carballo_Dacuy.docx")

    filas = cargar_benchmarks(csv_path)
    doc = Document()

    # Ajuste de margenes
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Membrete UADER FCyT
    logo = extraer_logo_uader(base_dir)
    if logo and os.path.exists(logo):
        doc.add_picture(logo, width=Inches(2.8))

    p_carrera = doc.add_paragraph()
    r1 = p_carrera.add_run("LIC. EN SISTEMAS DE INFORMACIÓN\nComputación Avanzada - 4° AÑO\n")
    r1.bold = True
    r1.italic = True
    r1.font.size = Pt(11)
    r2 = p_carrera.add_run("Alumnos: Carballo Tobías, Dacuy Santino\nDocente: Dra. María Fabiana Piccoli")
    r2.bold = True
    r2.font.size = Pt(11)

    titulo = doc.add_heading("Práctico Final Integrador: Generación de Cartoons a partir de Imágenes", level=0)
    for r in titulo.runs:
        r.font.color.rgb = RGBColor(10, 45, 95)

    # 1. Marco Teorico
    agregar_titulo_seccion(doc, "1. Marco Teórico: Computación Paralela Híbrida", nivel=1)
    doc.add_paragraph(
        "El procesamiento digital de imágenes de alta resolución requiere aplicar millones de operaciones "
        "aritméticas sobre matrices densas de píxeles. Ante los límites físicos de escalamiento de frecuencia "
        "de los procesadores secuenciales, las arquitecturas modernas de alto rendimiento combinan dos niveles "
        "jerárquicos de paralelismo: nodos multiprocesador de memoria compartida (multi-core SMP/NUMA) "
        "interconectados mediante redes de alta velocidad en esquemas de memoria distribuida."
    )
    doc.add_paragraph(
        "La Computación Paralela Híbrida integra ambos modelos dentro de una misma aplicación: utiliza el "
        "pasaje explícito de mensajes (MPI) para particionar el dominio global entre procesos independientes "
        "(paralelismo de grano grueso) y emplea directivas de memoria compartida (OpenMP) para distribuir "
        "los bucles de cálculo intensivo entre los hilos internos de cada proceso (paralelismo de grano fino)."
    )

    agregar_titulo_seccion(doc, "1.1. Características, Ventajas y Desventajas", nivel=2)
    items_ventajas = [
        "Reducción del overhead de red y memoria: Al lanzar menos procesos MPI (P procesos con T hilos cada uno, en lugar de P*T procesos MPI puros), disminuye la cantidad de mensajes colectivos y el número de fronteras (filas fantasma) duplicadas.",
        "Uso eficiente de la jerarquía de caché: Los T hilos de un mismo proceso MPI comparten estructuras de solo lectura (como las tablas LUT de posterizado y las filas fantasma ya recibidas), aprovechando las cachés L1/L2/L3.",
        "Soporte de hilos mediante MPI_THREAD_FUNNELED: En nuestra implementación híbrida inicializamos MPI con MPI_Init_thread(..., MPI_THREAD_FUNNELED, ...), donde el hilo maestro realiza todas las comunicaciones MPI fuera de las regiones paralelas de OpenMP, evitando contención interna en el driver de red.",
        "Desafíos del modelo híbrido: Requiere diseñar cuidadosamente dos niveles de particionado y sincronización, y elegir la relación óptima (P x T) acorde a la cantidad de núcleos físicos del procesador."
    ]
    for it in items_ventajas:
        doc.add_paragraph(it, style="List Bullet")

    # 2. Algoritmo Cartoon
    agregar_titulo_seccion(doc, "2. Características de la Aplicación: Algoritmo de Generación de Cartoons", nivel=1)
    doc.add_paragraph(
        "Una imagen digital RGB de dimensiones W x H píxeles se almacena como un arreglo unidimensional contiguo "
        "ordenado por filas (row-major order) de W * H * 3 bytes, donde cada píxel contiene las componentes "
        "(R, G, B) en el rango [0, 255]. Cumpliendo con el enunciado, cuando una operación de vecindad requiere "
        "acceder a coordenadas fuera de los bordes de la imagen, el valor neutro adoptado es negro (0, 0, 0)."
    )
    doc.add_paragraph(
        "Para transformar la imagen original en un Cartoon se aplican dos procesos independientes y finalmente "
        "se suman las imágenes resultantes:"
    )

    pasos_alg = [
        "Rama 1 — Detección de Bordes (3 pasos obligatorios, Pág. 57 del Apunte): "
        "(1) Filtrado / Borroneo mediante convolución con máscara pasa-bajo de media aritmética de 3x3 (peso 1/9) o 5x5 (peso 1/25) para eliminar el ruido; "
        "(2) Resaltado mediante aproximación del gradiente 2D aplicando los operadores direccionales de Sobel (Gx y Gy en 3x3 o 5x5) y calculando su magnitud |G| = sqrt(Gx^2 + Gy^2) ponderada por luminancia (Y = 0.30*R + 0.59*G + 0.11*B); "
        "(3) Detección mediante operación de umbralado (Umbralize, Pág. 10 y 58) que clasifica en blanco (255) los píxeles cuyo gradiente supera el umbral y en negro (0) el resto.",
        "Rama 2 — Posterizado por Canal de Color (Pág. 10 y 12 del Apunte): "
        "Divide el rango continuo [0, 255] de cada canal (R, G y B) en un número pequeño de rangos discretos (r = 3 o r = 9 valores) mediante una tabla de consulta (LUT) precalculada y suma los tres canales posterizados en una sola imagen RGB.",
        "Combinación Final: Se suman componente a componente con saturación en 255 la imagen resultante de la detección de bordes y la imagen posterizada: Cartoon(y,x,c) = min(255, Posterizado(y,x,c) + Bordes(y,x,c))."
    ]
    for p_txt in pasos_alg:
        doc.add_paragraph(p_txt, style="List Number")

    p_img1 = os.path.join(plots_dir, "panel_etapas_cartoon.png")
    if os.path.exists(p_img1):
        doc.add_picture(p_img1, width=Inches(6.5))
        cap = doc.add_paragraph("Figura 1: Etapas sucesivas del pipeline de generación de Cartoon sobre la escena de 800x800.")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_img2 = os.path.join(plots_dir, "comparativa_parametros_cartoon.png")
    if os.path.exists(p_img2):
        doc.add_picture(p_img2, width=Inches(5.4))
        cap = doc.add_paragraph("Figura 2: Comparativa visual variando tamaño de filtro (3x3 vs 5x5) y rangos de posterizado (3 vs 9).")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. Estrategias de Division y Comunicaciones
    agregar_titulo_seccion(doc, "3. Estrategias de División de Problemas y Comunicaciones", nivel=1)
    doc.add_paragraph(
        "Se implementaron cuatro versiones en lenguaje C (estándar C11) compartiendo el núcleo matemático "
        "(include/cartoon_core.h) para garantizar una comparación justa de Speedup y Eficiencia:"
    )
    estrategias = [
        "1. Versión Secuencial (src/cartoon_seq.c): Ejecuta secuencialmente el borroneo, resaltado y umbralado de bordes, posterizado y suma sobre arreglos contiguos 1D, estableciendo el tiempo de referencia Ts.",
        "2. Versión en Memoria Compartida con OpenMP (src/cartoon_omp.c): Paraleliza los bucles externos de filas mediante '#pragma omp parallel for default(none) schedule(static)'. Se eligió planificación estática porque el costo computacional por fila es idéntico en toda la imagen. Cada hilo escribe en filas disjuntas (sin condiciones de carrera) y la barrera implícita al finalizar el Paso 1 asegura que toda la imagen borroneada esté lista antes de que el Paso 2 lea filas vecinas.",
        "3. Versión en Memoria Distribuida con MPI (src/cartoon_mpi.c): Aplica descomposición de dominio 1D por bloques de filas equilibrados (filas[r] = H/P + (r < H%P)), distribuyendo la imagen con MPI_Scatterv y recolectando el resultado con MPI_Gatherv. Para resolver las dependencias espaciales de las convoluciones 3x3 (radio R=1) y 5x5 (radio R=2), cada proceso aloja m + 2R filas e intercambia R filas fantasma (Halo/Ghost Rows) con sus vecinos superior (rank-1) e inferior (rank+1) mediante MPI_Sendrecv antes del Borroneo y antes de Sobel. En los extremos (rank 0 y rank P-1) se utiliza MPI_PROC_NULL, conservando en 0 (negro) las filas fantasma exteriores.",
        "4. Versión Paralela Híbrida MPI + OpenMP (src/cartoon_hybrid.c): Combina el particionado inter-proceso e intercambio de filas fantasma de MPI con la paralelización intra-proceso de bucles mediante OpenMP, fusionando además los bucles de Sobel, Posterizado y Suma para maximizar el reuso de caché por fila."
    ]
    for est in estrategias:
        doc.add_paragraph(est, style="List Bullet")

    # 4. Resultados y Analisis de Desempenio
    agregar_titulo_seccion(doc, "4. Análisis de Desempeño: Tiempo, Speedup y Eficiencia (Inciso 5)", nivel=1)
    doc.add_paragraph(
        "Las pruebas se ejecutaron sobre un procesador AMD Ryzen 7 5700U (8 núcleos físicos / 16 hilos lógicos SMT) "
        "utilizando GCC 16.2.0 (-O3 -fopenmp) y Microsoft MPI v10.1. Previo al registro de tiempos se comprobó que "
        "las 4 versiones producen exactamente la misma imagen de salida píxel a píxel (Diferencia Máxima = 0)."
    )

    agregar_titulo_seccion(doc, "4.1. Tabla de Resultados — Imagen 800x800", nivel=2)
    agregar_tabla_resolucion(doc, filas, "800x800")

    agregar_titulo_seccion(doc, "4.2. Tabla de Resultados — Imagen 2000x2000", nivel=2)
    agregar_tabla_resolucion(doc, filas, "2000x2000")

    agregar_titulo_seccion(doc, "4.3. Tabla de Resultados — Imagen 5000x5000", nivel=2)
    agregar_tabla_resolucion(doc, filas, "5000x5000")

    agregar_titulo_seccion(doc, "4.4. Gráficos Comparativos y Discusión", nivel=2)
    for fig_file, caption in [
        ("openmp_speedup_eficiencia.png", "Figura 3: Speedup y Eficiencia en Memoria Compartida (OpenMP) para 2, 4, 8 y 16 threads."),
        ("mpi_speedup_eficiencia.png", "Figura 4: Speedup y Eficiencia en Memoria Distribuida (MPI) para 2, 4 y 8 procesos."),
        ("comparativa_paradigmas_5000.png", "Figura 5: Comparativa global entre Secuencial, OpenMP, MPI e Híbrido (Imagen 5000x5000, Filtro 5x5)."),
        ("impacto_filtros_y_comunicacion.png", "Figura 6: Impacto del tamaño de filtro (3x3 vs 5x5) y desglose Cómputo vs Comunicación MPI."),
    ]:
        p_fig = os.path.join(plots_dir, fig_file)
        if os.path.exists(p_fig):
            doc.add_picture(p_fig, width=Inches(6.4))
            c = doc.add_paragraph(caption)
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    discusion = [
        "Escalamiento con el tamaño de la imagen: Al pasar de 800x800 (640 mil píxeles) a 5000x5000 (25 millones de píxeles), la relación cómputo/comunicación crece sustancialmente. En 5000x5000 se alcanzan Speedups de hasta 8.07x en OpenMP (filtro 3x3) y 7.71x en la versión Híbrida 4Px4T (filtro 5x5, reduciendo el tiempo de 7.20 s a 0.93 s).",
        "Impacto del filtro (3x3 vs 5x5) y del posterizado (3 vs 9): El filtro de 5x5 requiere 25 accesos por píxel frente a 9 del filtro 3x3 (2.78x más operaciones teóricas), incrementando el tiempo de cómputo en ~2.57x y mejorando la eficiencia paralela. En cambio, variar el posterizado de 3 a 9 rangos no altera el tiempo de ejecución gracias al uso de la tabla LUT precalculada en O(1).",
        "Límite de núcleos físicos vs hilos lógicos (SMT): Hasta 8 workers (1 por núcleo físico del Ryzen 7 5700U), la eficiencia se mantiene elevada (73% a 94% en 5000x5000). Al escalar a 16 workers (usando los 16 hilos lógicos SMT), el Speedup continúa subiendo hasta 7.71x–8.07x, pero la eficiencia se sitúa entre 44% y 50% debido a que pares de hilos comparten las unidades de ejecución de cada núcleo físico."
    ]
    for d in discusion:
        doc.add_paragraph(d, style="List Bullet")

    # 5. Conclusiones
    agregar_titulo_seccion(doc, "5. Conclusiones", nivel=1)
    doc.add_paragraph(
        "La implementación de los cuatro modelos permitió verificar que el procesamiento de imágenes por "
        "convolución y transformación tonal presenta una excelente escalabilidad bajo descomposición de "
        "dominio por bloques de filas. En memoria compartida (OpenMP), el paralelismo se logra sin sobrecarga "
        "de copia de buffers, mientras que en memoria distribuida (MPI) el uso de filas fantasma (Halo Rows) "
        "reduce las comunicaciones a menos del 7% del tiempo total en imágenes de alta resolución. Finalmente, "
        "la arquitectura Híbrida (MPI + OpenMP) demostró ser la alternativa más rápida en el escenario de "
        "mayor carga (5000x5000 con filtro 5x5), combinando lo mejor de ambos paradigmas."
    )

    doc.save(out_docx)
    print(f"[OK] Informe Word generado exitosamente en: {out_docx}")


if __name__ == "__main__":
    main()
